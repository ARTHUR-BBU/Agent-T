# -*- coding: utf-8 -*-
"""B2-3 建议层考试阅卷器（考试1021 起固化）。

用法（仓库根目录）：
    CATEGORY_SUGGEST_ENABLED=true python -X utf8 -m tools.b23.grade_exam \
        --exam-dir "考试1021/最后一场考试_脱敏合同样本包" --run run1

纪律（spec-b23-trail-budget / 考试 README 同款）：
- 只喂 01_考试样本_喂给门童/ 下的卷面；02_考官答案_勿喂给门童/ 绝不作模型输入；
- 阅卷用生产同一份代码与提示词（app.services.category_suggest 单点）；
- 卷面文件所在目录须在 .git/info/exclude（脚本强制校验，防误提交）；
- 成绩册（json）只落考试目录，不进 git；报告不引用合同原文。
"""
from __future__ import annotations

import argparse
import csv
import json
import pathlib
import subprocess
import time

# 期望品类（考官中文口径）→ 引擎八类 id
_ZH_TO_ID = {"采购": "procurement", "租赁": "lease", "保密协议": "nda",
             "技术服务": "technical_service"}
# 四收类：判成这些=收进审查；判成其余（consulting/labor_dispatch/tech_transfer/other）=不收
_SUPPORTED = {"procurement", "lease", "nda", "technical_service"}
# 拒收话术只在建议命中显式拒类时出现（other 走存疑放行，语义不同）
_EXPLICIT_REJECT = {"consulting", "labor_dispatch", "tech_transfer"}


def _ensure_ignored(exam_dir: pathlib.Path) -> None:
    """考试目录必须在 git 排除清单（防卷面/成绩册误提交）。"""
    probe = exam_dir / ".grade_probe"
    probe.touch()
    try:
        r = subprocess.run(["git", "check-ignore", "-q", str(probe)])
        if r.returncode != 0:
            raise SystemExit(
                f"FAIL: {exam_dir} 未被 git 忽略——先加入 .git/info/exclude 再阅卷")
    finally:
        probe.unlink()


def _docx_text(path: pathlib.Path) -> str:
    import docx  # python-docx（生产解析同款库）
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:  # 卷面含清单/SLA 表格（考官明示的信号载体）
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def judge(pred, gold_conclusion: str, gold_id, performed: bool = True) -> tuple[bool, str]:
    """单题判分（纯函数，供单测）。

    收类：判中金标品类才算对。拒收（Codex 勘正 2026-10-22）：只有显式拒类
    （consulting/labor_dispatch/tech_transfer）才算拒收命中——other 是「存疑
    按用户原选择继续」不是拒收；None（未判/基础设施故障）更不是。"""
    if gold_conclusion == "收":
        ok = (pred == gold_id)
        verdict = ("对" if ok else
                   "误收" if pred in _SUPPORTED else
                   "误拒" if pred in _EXPLICIT_REJECT else "判other")
    else:
        ok = bool(performed and pred in _EXPLICIT_REJECT)
        verdict = ("对（显式拒类）" if ok else
                   f"未通过（{'未判' if not performed else pred or 'other 存疑放行'}）")
    return ok, verdict


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--exam-dir", required=True)
    ap.add_argument("--run", default="run1")
    ap.add_argument("--selected", default="procurement",
                    help="考生视角用户所选品类（真实上传流同款）")
    args = ap.parse_args()

    exam_dir = pathlib.Path(args.exam_dir)
    papers_dir = next(d for d in exam_dir.iterdir() if d.name.startswith("01_"))
    answer_dir = next(d for d in exam_dir.iterdir() if d.name.startswith("02_"))
    _ensure_ignored(exam_dir)

    csv_path = next(answer_dir.glob("*.csv"))
    raw_rows = list(csv.DictReader(csv_path.read_text(encoding="utf-8-sig").splitlines()))
    # 双 schema 兼容：高仿卷用 文件名/期望结论/期望品类；真实卷用
    # 样本编号/期望品类(含「野样本·X(应拒收)」)/是否野样本(应拒收)。
    answers = {}
    for row in raw_rows:
        if "文件名" in row:  # 高仿卷 schema
            answers[row["文件名"]] = {
                "conclusion": row["期望结论"].strip(),
                "gold_zh": row["期望品类"].strip(),
                "difficulty": row.get("难度", ""),
                "trap": row.get("设计意图/易错点", ""),
            }
        else:  # 真实卷 schema（溯源清单）
            is_wild = row.get("是否野样本(应拒收)", "").strip() == "是"
            answers[f"{row['样本编号'].strip()}.docx"] = {
                "title": row.get("原合同名称", ""),
                "conclusion": "拒收" if is_wild else "收",
                "gold_zh": row["期望品类"].strip(),
                "difficulty": "边界样本" if row.get("是否边界样本", "").strip() == "是" else "常规",
                "trap": row.get("难度与备注", ""),
            }

    from app.services import category_suggest as cs
    if not cs.is_enabled():
        raise SystemExit("FAIL: CATEGORY_SUGGEST_ENABLED 未开——阅卷需真实调用")

    rows = []
    papers = sorted(papers_dir.glob("*.docx"))
    t0 = time.time()
    for paper in papers:
        ans = answers[paper.name]
        gold_conclusion = ans["conclusion"]                # 收 | 拒收
        gold_zh = ans["gold_zh"]
        gold_id = _ZH_TO_ID.get(gold_zh)                   # 拒收题为 None
        text = _docx_text(paper)
        out = cs.suggest_category(text, args.selected)
        pred = out.result.category if (out.performed and out.result) else None
        conf = out.result.confidence if out.result else None
        tier = cs.confidence_tier(conf) if conf is not None else None
        if gold_conclusion == "收":
            ok = (pred == gold_id)
            verdict = ("对" if ok else
                       "误收" if pred in _SUPPORTED else
                       "误拒" if pred in _EXPLICIT_REJECT else "判other")
        else:
            # 假绿防线（真实卷 run2 教训 + Codex 勘正）：None 与 other 都不算拒收命中
            ok, verdict = judge(pred, gold_conclusion, gold_id, performed=out.performed)
        rows.append({
            "paper": paper.name, "title": ans.get("合同标题") or ans.get("title", ""),
            "gold_conclusion": gold_conclusion, "gold": gold_id,
            "pred": pred, "conf": round(conf, 4) if conf else None,
            "tier": tier, "ok": ok, "verdict": verdict,
            "difficulty": ans["difficulty"], "trap": ans["trap"],
            "skip": out.skip_reason,
        })
        time.sleep(0.3)

    elapsed = round(time.time() - t0, 1)
    out_path = exam_dir / f"results-final-{args.run}.json"
    out_path.write_text(json.dumps(
        {"run": args.run, "ts": time.strftime("%Y-%m-%d %H:%M"),
         "selected": args.selected, "elapsed_s": elapsed, "rows": rows},
        ensure_ascii=False, indent=1), encoding="utf-8")

    collect = [r for r in rows if r["gold_conclusion"] == "收"]
    reject = [r for r in rows if r["gold_conclusion"] == "拒收"]
    print(f"=== 最后一场考试（{args.run}，{len(rows)} 份，{elapsed}s）===")
    for r in rows:
        print(("✅" if r["ok"] else "❌"), r["paper"], r["title"],
              f"金标={r['gold_conclusion']}/{r['gold']}", f"判={r['pred']}",
              f"置信={r['conf']}", r["verdict"])
    print(f"\n收类判断：{sum(r['ok'] for r in collect)}/{len(collect)}")
    print(f"拒收题拦截：{sum(r['ok'] for r in reject)}/{len(reject)}")
    wrong = [r for r in rows if not r["ok"]]
    if wrong:
        print("错题清单：")
        for r in wrong:
            print(" ", r["paper"], r["title"], "→", r["pred"],
                  f"(置信 {r['conf']}, 档 {r['tier']})")
    print("成绩册：", out_path)


if __name__ == "__main__":
    main()
