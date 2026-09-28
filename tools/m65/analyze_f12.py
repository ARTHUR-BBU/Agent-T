"""F-1/F-2 靶向分析器：拿本地源文件当标准答案，判定审查结果的真缺/漏报/误报。

用法：python -X utf8 tools/m65/analyze_f12.py --run docs/m65/run2.json [--run docs/m65/run1.json]

F-1（漏报）：清单「期限/价款与支付/标的/保密/知识产权」报「未找到」时，
  用本地 docx 段落+表格全文关键词回归——文中明明有对应措辞 = 漏报嫌疑。
F-2（误报）：需关注 note 含「单方免除」时，查本地全文该结论引用的条款
  上下文是否含「任何一方/双方」对等表述——含 = 误报嫌疑。
输出逐条判定表（供老钱裁决金标引用），不判死刑只标嫌疑，最终裁决权在老钱。
"""
from __future__ import annotations

import argparse
import io
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]

# F-1：清单项 → 本地全文「存在即漏报嫌疑」的措辞信号（多信号命中更可信）
_F1_SIGNALS: dict[str, tuple[str, ...]] = {
    "期限": ("合同期限", "租赁期限", "供货服务期", "履行期限", "交付期限", "服务期限",
             "租赁期", "工期", "有效期"),
    "价款与支付": ("支付", "付款", "价款", "货款", "报酬", "服务费", "租金", "结算"),
    "标的": ("标的", "材料名称", "品名", "设备名称", "货物名称", "租赁物资", "承揽项目"),
    "保密": ("保密", "商业秘密", "机密"),
    "知识产权": ("知识产权", "著作权", "专利", "版权"),
}

# F-2：对等免责的对立信号——出现即倾向「误报嫌疑」
_F2_EQUAL_SIGNALS = ("任何一方", "双方", "各自")


def _full_text(fixtures: Path, fname: str) -> str:
    """段落+表格全文（表格 cell 去重拼接，与审查端解析口径近似）。"""
    from docx import Document
    d = Document(str(fixtures / fname))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            uniq = list(dict.fromkeys(cells))
            if any(uniq):
                parts.append(" | ".join(uniq))
    return "\n".join(parts)


def analyze_f1(rec: dict, text: str) -> list[dict]:
    out: list[dict] = []
    for i in rec.get("items_detail") or []:
        if i.get("status") != "未找到":
            continue
        name = str(i.get("name") or "")
        signals = None
        for key, sigs in _F1_SIGNALS.items():
            if key in name:
                signals = sigs
                break
        if signals is None:
            continue
        hits = [s for s in signals if s in text]
        out.append({
            "type": "F1", "file": rec.get("file"), "item": name,
            "verdict": "漏报嫌疑" if hits else "真缺失（倾向）",
            "signal_hits": hits[:6],
        })
    return out


def analyze_f2(rec: dict, text: str) -> list[dict]:
    out: list[dict] = []
    for i in rec.get("items_detail") or []:
        note = str(i.get("note") or "")
        if i.get("status") == "需关注" and "单方免除" in note:
            # 找到该结论引用的原文（quote 截断过，退回全文检索对等信号）
            quote = str(i.get("quote") or "")
            anchor = quote.strip("…。 ")[:20]
            ctx = text
            if anchor and anchor in text:
                pos = text.find(anchor)
                ctx = text[max(0, pos - 300): pos + 300]
            equal_hits = [s for s in _F2_EQUAL_SIGNALS if s in ctx]
            out.append({
                "type": "F2", "file": rec.get("file"), "item": i.get("name"),
                "verdict": "误报嫌疑（原文对等）" if equal_hits else "待人工复核",
                "signal_hits": equal_hits, "quote_head": quote[:60],
            })
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", action="append", required=True, help="run json（可多次）")
    ap.add_argument("--fixtures", default="fixtures-real")
    ap.add_argument("--out", default="")
    args = ap.parse_args()
    fixtures = ROOT / args.fixtures

    rows: list[dict] = []
    for run_path in args.run:
        d = json.loads(io.open(run_path, encoding="utf-8").read())
        for group in ("contracts", "contracts_extra"):
            for rec in d.get(group) or []:
                if rec.get("status") != "done":
                    continue
                fname = rec.get("file") or ""
                try:
                    text = _full_text(fixtures, fname)
                except Exception as exc:  # noqa: BLE001
                    print(f"[warn] {fname} 本地解析失败：{exc}", file=sys.stderr)
                    continue
                rows.extend(analyze_f1(rec, text))
                rows.extend(analyze_f2(rec, text))

    for r in rows:
        sig = "、".join(r["signal_hits"]) or "—"
        print(f"[{r['type']}] {r['file'][:44]:46} | {r['item']} | {r['verdict']} | 信号: {sig}")
    f1 = [r for r in rows if r["type"] == "F1"]
    f2 = [r for r in rows if r["type"] == "F2"]
    print(f"\n汇总：F-1 判定 {len(f1)} 条（漏报嫌疑 {sum(1 for r in f1 if '漏报' in r['verdict'])}）；"
          f"F-2 判定 {len(f2)} 条（误报嫌疑 {sum(1 for r in f2 if '误报' in r['verdict'])}）")
    if args.out:
        io.open(args.out, "w", encoding="utf-8").write(
            json.dumps(rows, ensure_ascii=False, indent=2))
        print(f"saved -> {args.out}")


if __name__ == "__main__":
    main()
