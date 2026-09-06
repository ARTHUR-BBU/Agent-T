"""M4 审查报告导出：把 store 里的既有审查结果拼装成 docx。

红线（与 M3.5 评分卡架构同源）：
- 报告只做汇总与展示，不产生任何新判断——档位一律照抄规则引擎结果；
- 补盲候选单独成节，必须标注「模型候选，需人工确认」；
- 免责句是必备项：scorecard 没带时由代码补默认句；
- 合同全文不写入报告（只含原文摘句），避免报告本身成为泄露面。
"""
from __future__ import annotations

import io
import re
from typing import Any

from app.services.scorecard import scrub_forbidden

# 免责句兜底：评分卡未生成（无 Key / 降级）时报告也必须带免责声明
DEFAULT_DISCLAIMER = (
    "本报告由系统自动生成，仅供签约前自查参考，不构成法律意见。"
    "重要合同请在签署前交由律师或法务人员审查。"
)

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_ATTENTION = "需关注"
_NOT_FOUND = "未找到"
_NA = "本类不适用"

# python-docx(lxml) 遇控制字符直接抛错：合同原文经 errors='replace'/pypdf/docling
# 提取后可能残留 \x0b\x0c 等，不清洗则该条审查的报告导出永久 500（肉饼审计 P1）
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _sanitize(value: Any) -> Any:
    """递归清洗字符串中的控制字符（单点收口：所有入 docx 的数据都过这里）."""
    if isinstance(value, str):
        return _CONTROL_CHARS.sub("", value)
    if isinstance(value, dict):
        return {k: _sanitize(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_sanitize(v) for v in value]
    return value


def _scrub(text: Any) -> str:
    """自由文本字段的禁语二次清洗（纵深防御，肉饼/小智娘 M4 审计共同建议）。

    生成层（scorecard/blind_spot）已有一道清洗；报告是印给客户看的文书，
    这里对模型的自由文本字段再过一遍——确定性替换，不是新判断，红线1不受影响。
    """
    return scrub_forbidden(str(text or ""))


def build_report_docx(row: dict[str, Any]) -> bytes:
    """从 store 行拼装报告，返回 docx 字节流。纯展示，无 LLM 调用。"""
    from docx import Document

    row = _sanitize(row)

    doc = Document()
    _cover(doc, row)
    _conclusion(doc, row)
    _attention_table(doc, row)
    _item_details(doc, row)
    _blind_candidates(doc, row)
    _policy_quotes(doc, row)
    _appendix(doc, row)
    _footer(doc, row)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- 各节 ----------

def _cover(doc: Any, row: dict[str, Any]) -> None:
    doc.add_heading("合同审查报告", level=0)
    rows = [
        ("合同文件", row.get("filename") or "未命名"),
        ("合同类型", row.get("category_label") or row.get("category") or ""),
        ("审查时间", row.get("created_at") or "未记录"),
        ("报告编号", row.get("id") or ""),
    ]
    for label, value in rows:
        p = doc.add_paragraph()
        run = p.add_run(f"{label}：")
        run.bold = True
        p.add_run(str(value))


def _known_statuses() -> tuple[str, ...]:
    return ("通过", _ATTENTION, _NOT_FOUND, _NA)


def _conclusion(doc: Any, row: dict[str, Any]) -> None:
    doc.add_heading("一、审查结论汇总", level=1)
    items = row.get("items") or []
    counts = {s: 0 for s in _known_statuses()}
    unknown = 0
    for it in items:
        status = it.get("status")
        # isinstance 守卫（小智娘 P2-1）：list/dict 等不可哈希脏值走从严分支，
        # 不能在 dict 成员判断上 TypeError 炸 500——fail-closed 而非 fail-loud
        if isinstance(status, str) and status in counts:
            counts[status] += 1
        else:
            # 未知/变体档位不许静默吞掉（遗留项①，肉饼 P2-2）：计数守恒，
            # 且从严往风险侧倒，绝不落进「全部通过」的错误背书
            unknown += 1
    doc.add_paragraph(
        f"共核查 {len(items)} 项：通过 {counts['通过']} · 需关注 {counts[_ATTENTION]}"
        f" · 未找到 {counts[_NOT_FOUND]} · 本类不适用 {counts[_NA]}"
        + (f" · 无法识别档位 {unknown} 项" if unknown else "")
    )
    if unknown:
        doc.add_paragraph("上述无法识别档位已按「需关注」从严处理，请人工复核。")

    sc = row.get("scorecard") or {}
    if sc.get("available") and isinstance(sc.get("total"), (int, float)):
        tier = sc.get("tier") or {}
        label = tier.get("label") or ""
        doc.add_paragraph(f"参考评分：{sc['total']}/100（{label}）" if label else f"参考评分：{sc['total']}/100")
        if sc.get("summary"):
            doc.add_paragraph(_scrub(sc["summary"]))
        for cap in sc.get("caps_applied") or []:
            doc.add_paragraph(str(cap), style="List Bullet")
        doc.add_paragraph("评分由模型生成、仅供排序参考，各项结论以逐条规则结果为准。")
    else:
        reason = sc.get("reason") or "未生成"
        doc.add_paragraph(f"参考评分未生成（{reason}），各项结论以逐条规则结果为准。")


def _attention_table(doc: Any, row: dict[str, Any]) -> None:
    doc.add_heading("二、需关注与未找到汇总", level=1)
    known = _known_statuses()
    hard = [
        it for it in (row.get("items") or [])
        if it.get("status") in (_ATTENTION, _NOT_FOUND) or it.get("status") not in known
    ]
    if not hard:
        doc.add_paragraph("无——全部适用项均通过。")
        return
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    for i, text in enumerate(("条目", "档位", "说明")):
        header[i].text = text
    for it in hard:
        cells = table.add_row().cells
        cells[0].text = it.get("name") or ""
        status = it.get("status") or ""
        # 未知档位进表时显式标注从严口径，不让读者误以为档位可信
        cells[1].text = status if status in (_ATTENTION, _NOT_FOUND) else f"{status or '空'}（按需关注处理）"
        cells[2].text = _scrub(it.get("note"))


def _item_details(doc: Any, row: dict[str, Any]) -> None:
    doc.add_heading("三、逐条明细", level=1)
    shown = [
        it for it in (row.get("items") or [])
        if it.get("status") != _NA
    ]
    if not shown:
        doc.add_paragraph("无适用清单项。")
        return
    for it in shown:
        doc.add_heading(f"【{it.get('status') or '未判定'}】{it.get('name') or ''}", level=2)
        doc.add_paragraph(_scrub(it.get("note")) or "（无说明）")
        quote = it.get("quote") or ""
        p = doc.add_paragraph()
        run = p.add_run("原文摘句：")
        run.bold = True
        p.add_run(quote if quote.strip() else "暂无")


def _blind_candidates(doc: Any, row: dict[str, Any]) -> None:
    candidates = row.get("blind_candidates") or []
    if not candidates:
        return
    doc.add_heading("四、模型补盲候选（需人工确认）", level=1)
    doc.add_paragraph(
        "以下为模型提出的候选风险，未经规则引擎确认，不构成审查结论，"
        "请人工核实后再决定是否采信。"
    )
    for it in candidates:
        name = it.get("name") or ""
        marker = "· 来自评分卡点名" if it.get("named_by_scorecard") else ""
        doc.add_heading(f"候选｜{name}{marker}", level=2)
        doc.add_paragraph(_scrub(it.get("note")) or "（无说明）")
        quote = it.get("quote") or ""
        p = doc.add_paragraph()
        run = p.add_run("原文摘句：")
        run.bold = True
        p.add_run(quote if quote.strip() else "暂无")


def _policy_quotes(doc: Any, row: dict[str, Any]) -> None:
    doc.add_heading("五、政策摘句", level=1)
    policies = [p for p in (row.get("policies") or []) if str(p).strip()]
    if not policies:
        doc.add_paragraph("本次审查未引用政策条款。")
        return
    for policy in policies:
        doc.add_paragraph(_scrub(policy), style="List Bullet")


def _appendix(doc: Any, row: dict[str, Any]) -> None:
    doc.add_heading("附注", level=1)
    na_names = [
        it.get("name") or "" for it in (row.get("items") or [])
        if it.get("status") == _NA
    ]
    if na_names:
        doc.add_paragraph("本类不适用（未核查）：" + "、".join(n for n in na_names if n))
    skips = [s for s in (row.get("blind_skipped_messages") or []) if str(s).strip()]
    if skips:
        doc.add_paragraph("补盲候选跳过：" + "；".join(str(s) for s in skips))


def _footer(doc: Any, row: dict[str, Any]) -> None:
    sc = row.get("scorecard") or {}
    disclaimer = _scrub(sc.get("disclaimer")).strip() or DEFAULT_DISCLAIMER
    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(disclaimer)
    run.italic = True
    doc.add_paragraph(f"报告编号 {row.get('id') or ''} · 由合同审查 Agent 自动生成")
