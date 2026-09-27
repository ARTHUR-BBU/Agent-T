"""批 2c 回归：决定记录（轻量版）+ 报告可追溯 + 真幂等 + 一致性检测。

实现稿 v1.5（docs/evidence-batch2c-impl.md）§6 测试矩阵。
红线：变异验证回滚必须变红——删指纹快照/删一致性检测/删 adopt 前置校验/
删锁/报告绕过视图，由本文件与门禁人工变异共同承担。
"""
from __future__ import annotations

import copy
import io
import json

import pytest

from app.services import llm_ask as llm_ask_service  # noqa: F401 统一导入风格
from app.services.evidence import build_evidence, document_version_for
from app.services.store import store as store_module
from tests.test_evidence_claims import client  # noqa: F401 复用模块级 TestClient

_TEXT = (
    "第一条 乙方需在合同签订后20日内完成送货、安装调试，交付甲方使用。\n"
    "第二条 乙方保证货物为正品，若出现质量问题，乙方可协助处理，甲方不得追究乙方违约责任。\n"
)
_DV = document_version_for(_TEXT)
_ITEM_QUOTE = "乙方需在合同签订后20日内完成送货、安装调试，交付甲方使用"
_OBJ_QUOTE = "乙方保证货物为正品，若出现质量问题，乙方可协助处理，甲方不得追究乙方违约责任"


def _ev(quote: str):
    return build_evidence(text=_TEXT, quote=quote, parse_source="rules", document_version=_DV)


def _row(**overrides) -> dict:
    row = {
        "filename": "c2c.txt", "category": "procurement", "status": "done",
        "stage": "done", "created_at": "2026-09-27 10:00",
        "scorecard": {}, "blind_candidates": [], "blind_skipped_messages": [],
        "blind_skipped_reason": None, "blind_enabled": False,
        "text": _TEXT, "document_version": _DV, "policies": [], "error": None,
        "items": [
            {"id": "delivery", "name": "交付时间", "status": "需关注",
             "note": "先款后货风险", "quote": _ITEM_QUOTE,
             "evidence": _ev(_ITEM_QUOTE)},
        ],
        "verify": {"available": True, "reason": None, "questions": [_question()]},
        "objections": {"available": True, "objections": [_objection()]},
    }
    row.update(overrides)
    return row


def _question(**overrides) -> dict:
    q = {
        "id": "q1", "source": "quality_obs", "source_subject_key": "k1",
        "question": "先款后货是否构成验收风险？", "title": "付款与验收顺序",
        "quote": _ITEM_QUOTE,
        "evidence": _ev(_ITEM_QUOTE),
        "verification": "verified", "status": "pending",
    }
    q.update(overrides)
    return q


def _objection(**overrides) -> dict:
    ob = {
        "item_id": "delivery", "rule_id": "delivery#r0", "rule_class": "heuristic",
        "direction": "false_positive", "quote": _ITEM_QUOTE,
        "counter_evidence": _OBJ_QUOTE,
        "counter_evidence_status": "present",
        "counter_evidence_ref": build_evidence(
            text=_TEXT, quote=_OBJ_QUOTE, parse_source="objection", document_version=_DV),
        "legal_reasoning": "二十日送货期属于正常履约安排，不构成先款后货的实质风险，需结合交付验收条款综合判断。",
        "stance_check": "与立场无关",
        "proposal": "为交付词表增加验收语境 unless",
        "accepted": True, "reject_reason": None,
        "clause_id": "c01", "clause_ambiguous": False,
        "evidence": _ev(_ITEM_QUOTE),
        "adopted": False, "needs_confirm": True,
    }
    ob.update(overrides)
    return ob


def _mk(rid_suffix: str = "", **overrides) -> str:
    return store_module.create(**_row(**overrides))


def _doc_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def _confirm(rid: str, **kw):
    body = {"question_id": "q1", "choice": "confirm", "human_note": "人工核对无误", "revised_quote": ""}
    body.update(kw)
    return client.post(f"/api/review/{rid}/confirm", json=body)


def _get_decision(rid: str) -> dict | None:
    body = client.get(f"/api/review/{rid}").json()
    qs = ((body.get("verify") or {}).get("questions") or [])
    return next((q.get("decision") for q in qs if q.get("id") == "q1"), None)


# ---------- A 组：决定写入 / 幂等 / 一致性 ----------

def test_t_a1_confirm_writes_decision_with_snapshot(monkeypatch):
    """T-A1：confirm → 11 键 store 记录 + 13 键 API 形状；快照与读路径派生一致。"""
    rid = _mk()
    r = _confirm(rid)
    assert r.status_code == 200 and r.json()["ok"] is True
    d = _get_decision(rid)
    assert d is not None
    store_keys = {k for k in d if k not in ("consistency", "consistency_reasons")}
    assert store_keys == {
        "decision_id", "decision_type", "actor", "authority", "claim_id",
        "claim_content_hash", "evidence_ids", "choice", "human_note",
        "revised_quote", "decided_at",
    }, store_keys
    assert d["decision_type"] == "human_confirm" and d["choice"] == "confirm"
    assert d["claim_id"].startswith("cl-") and d["claim_content_hash"].startswith("cc-")
    assert d["evidence_ids"] and all(e.startswith("ev-") for e in d["evidence_ids"])
    assert d["consistency"] == "consistent" and d["consistency_reasons"] == []


def test_t_a2_choice_change_new_id(monkeypatch):
    """T-A2：confirm→dispute = 新决定、记录整体替换。"""
    rid = _mk()
    _confirm(rid)
    d1 = _get_decision(rid)
    r = _confirm(rid, choice="dispute", human_note="改判：需复核验收条款")
    assert r.status_code == 200
    d2 = _get_decision(rid)
    assert d2["decision_id"] != d1["decision_id"]
    assert d2["choice"] == "dispute" and d2["decision_type"] == "human_dispute"


def test_t_a3_pending_confirm_disclosed_in_report(monkeypatch):
    """T-A3（Q1）：pending 问题确认 → 状态更新、无决定；pending-only 合同报告
    仍出现未编号披露段（v1.5：披露有了安放处）。"""
    pending_only = _question(id="q1", source="pending", evidence=_ev(_ITEM_QUOTE))
    rid = _mk(verify={"available": True, "reason": None, "questions": [pending_only]})
    r = _confirm(rid)
    assert r.status_code == 200 and r.json()["ok"] is True
    assert _get_decision(rid) is None, "pending 不写决定"
    body = client.get(f"/api/review/{rid}").json()
    q = body["verify"]["questions"][0]
    assert q["status"] == "confirmed" and q["decision"] is None
    data = client.get(f"/api/review/{rid}/report").content
    text = _doc_text(data)
    assert "以下问题已人工确认" in text and "未纳入决定链" in text


def test_t_a4_claim_drift_detected(monkeypatch):
    """T-A4：决定后改**真参与指纹的字段**（question 文本）→ claim_drift；
    三铁律：不删决定、不静默改写、不冒充已核实。"""
    rid = _mk()
    _confirm(rid)
    d1 = _get_decision(rid)
    # 改问题文本（verify_question 指纹白名单 = question/title）
    stored = store_module.get(rid)
    for q in stored["verify"]["questions"]:
        if q.get("id") == "q1":
            q["question"] = "改后的风险描述：交货期限是否过紧？"
    store_module.update(rid, verify=stored["verify"])
    d2 = _get_decision(rid)
    assert d2["decision_id"] == d1["decision_id"], "漂移不得删除或改写决定"
    assert "claim_drift" in d2["consistency_reasons"]
    assert d2["consistency"] == "degraded"
    body = client.get(f"/api/review/{rid}").json()
    assert "已核验" not in json.dumps(
        [q.get("decision") for q in body["verify"]["questions"]], ensure_ascii=False)


def test_t_a5_evidence_broken_detected(monkeypatch):
    """T-A5：决定后证据降级 → evidence_broken。"""
    rid = _mk()
    _confirm(rid)
    stored = store_module.get(rid)
    for q in stored["verify"]["questions"]:
        if q.get("id") == "q1":
            q["evidence"]["verification"] = "missing"
            q["evidence"]["evidence_id"] = ""
    store_module.update(rid, verify=stored["verify"])
    d = _get_decision(rid)
    assert "evidence_broken" in d["consistency_reasons"]


def test_t_a5b_both_reasons_coexist(monkeypatch):
    """T-A5b：漂移+失效并存 → reasons 两项并存、固定排序。"""
    rid = _mk()
    _confirm(rid)
    stored = store_module.get(rid)
    for q in stored["verify"]["questions"]:
        if q.get("id") == "q1":
            q["question"] = "并存场景：问题文本已改"
            q["evidence"]["verification"] = "missing"
            q["evidence"]["evidence_id"] = ""
    store_module.update(rid, verify=stored["verify"])
    d = _get_decision(rid)
    assert d["consistency_reasons"] == ["claim_drift", "evidence_broken"]


def test_t_a6_true_idempotency(monkeypatch):
    """T-A6（P2）：①完全相同重复提交字节稳定且 decided_at 不刷新；
    ②仅 note 变 → 说明修订（id 不变）；③choice 变 → 新 id。"""
    rid = _mk()
    _confirm(rid)
    d1 = _get_decision(rid)
    _confirm(rid)  # 完全相同
    d2 = _get_decision(rid)
    assert d2 == d1, "完全相同重复提交必须字节稳定（含 decided_at）"
    _confirm(rid, human_note="补充：已对照第二条")  # 说明修订
    d3 = _get_decision(rid)
    assert d3["decision_id"] == d1["decision_id"]
    assert d3["human_note"] == "补充：已对照第二条"
    assert d3["decided_at"] >= d1["decided_at"]
    assert d3["claim_id"] == d1["claim_id"]
    _confirm(rid, choice="dispute", human_note="改判")  # choice 变
    d4 = _get_decision(rid)
    assert d4["decision_id"] != d1["decision_id"] and d4["choice"] == "dispute"


# ---------- B 组：adopt ----------

def test_t_b1_adopt_decision_covers_all_refs(monkeypatch):
    """T-B1：adopt → objection_adopt；evidence_ids = 全部合格 refs 去重升序
    （primary/rebuts/counter 全覆盖，不漏 rebuts）。"""
    rid = _mk()
    r = client.post(f"/api/review/{rid}/objections/adopt", json={"index": 0})
    assert r.status_code == 200
    body = r.json()
    ob = body["objections"][0]
    assert ob["adopted"] is True
    d = ob["decision"]
    assert d["decision_type"] == "objection_adopt" and d["choice"] == "adopted"
    assert d["claim_id"].startswith("cl-")
    expected = sorted({r_["evidence_id"] for r_ in ob["evidence_refs"]})
    assert d["evidence_ids"] == expected
    relations = {r_["relation"] for r_ in ob["evidence_refs"]}
    assert {"primary", "rebuts", "counter"} <= relations, "样例必须覆盖三种关系"


def test_t_b2_adopt_without_claim_rejected(monkeypatch):
    """T-B2（审计 P1-2）：无编号异议采纳 → 422 拒绝，adopted 不变、无假状态。"""
    no_ev = _objection(evidence={"document_version": _DV, "quote": "", "start": None,
                                 "end": None, "clause_id": None,
                                 "verification": "missing", "parse_source": "objection",
                                 "evidence_id": ""})
    rid = _mk(objections={"available": True, "objections": [no_ev]})
    r = client.post(f"/api/review/{rid}/objections/adopt", json={"index": 0})
    assert r.status_code == 422
    stored = store_module.get(rid)
    assert stored["objections"]["objections"][0]["adopted"] is False
    assert stored["objections"]["objections"][0].get("decision") is None


def test_t_b3_adopt_response_annotated(monkeypatch):
    """T-B3（§2.5-④ 自查项）：adopt 响应不再是无标注原件——含派生字段。"""
    rid = _mk()
    r = client.post(f"/api/review/{rid}/objections/adopt", json={"index": 0})
    body = r.json()
    ob = body["objections"][0]
    assert ob["claim_id"].startswith("cl-"), "响应必须经流水线标注"
    assert ob["decision"] is not None and "consistency" in ob["decision"]


# ---------- C 组：报告 ----------

def test_t_c2_old_row_deterministic_and_annotated(monkeypatch):
    """T-C2（新口径）：旧档案导出含派生主张信息 + 两次导出字节一致。"""
    row = _row()
    row["items"][0]["evidence"] = {
        "document_version": _DV, "quote": _ITEM_QUOTE, "start": None, "end": None,
        "clause_id": None, "verification": "missing",
        "parse_source": "rules", "evidence_id": "",
    }  # 坏票 → 未编号 + 未能定位 两个降级分支
    row["verify"] = {"available": False, "reason": "not_attempted", "questions": []}
    row["objections"] = None
    rid = store_module.create(**row)
    data1 = client.get(f"/api/review/{rid}/report").content
    data2 = client.get(f"/api/review/{rid}/report").content
    assert data1 == data2, "确定性契约：同档案两次导出字节一致"
    text = _doc_text(data1)
    assert "未编号（证据不合格，未纳入证据链）" in text
    assert "未能定位到原文" in text


def test_t_c4_report_with_decision(monkeypatch):
    """T-C4：带决定导出 → 决定节 + 主张编号 + 证据引用列出现。"""
    rid = _mk()
    _confirm(rid)
    data = client.get(f"/api/review/{rid}/report").content
    text = _doc_text(data)
    assert "四、人工决定与确认" in text
    assert "cl-" in text and "确认" in text
    d = _get_decision(rid)
    assert d["evidence_ids"][0] in text, "证据引用列必须出现（Codex P1 落实）"


def test_t_c5_drift_report(monkeypatch):
    """T-C5：漂移后导出 → 报告出现中文降级文案，不冒充已核实。"""
    rid = _mk()
    _confirm(rid)
    stored = store_module.get(rid)
    for q in stored["verify"]["questions"]:
        if q.get("id") == "q1":
            q["question"] = "导出场景：问题文本已改"
    store_module.update(rid, verify=stored["verify"])
    text = _doc_text(client.get(f"/api/review/{rid}/report").content)
    assert "主张身份或内容已变化" in text


def test_t_c6_broken_report(monkeypatch):
    """T-C6：证据失效后导出 → 报告出现失效文案。"""
    rid = _mk()
    _confirm(rid)
    stored = store_module.get(rid)
    for q in stored["verify"]["questions"]:
        if q.get("id") == "q1":
            q["evidence"]["verification"] = "missing"
            q["evidence"]["evidence_id"] = ""
    store_module.update(rid, verify=stored["verify"])
    text = _doc_text(client.get(f"/api/review/{rid}/report").content)
    assert "决定引用的证据票据已失效或无法定位" in text


# ---------- D 组：API 往返 / 流水线同源 / 指纹免疫 ----------

def test_t_d1_verify_response_decision_13_keys(monkeypatch):
    """T-D1（P1-A 挂载点①）：verify 响应 decision 13 字段逐一存活，剥一键即红。"""
    rid = _mk()
    _confirm(rid)
    body = client.get(f"/api/review/{rid}/verify").json()
    d = next(q.get("decision") for q in body["questions"] if q["id"] == "q1")
    assert d is not None
    expected = {
        "decision_id", "decision_type", "actor", "authority", "claim_id",
        "claim_content_hash", "evidence_ids", "choice", "human_note",
        "revised_quote", "decided_at", "consistency", "consistency_reasons",
    }
    assert set(d) == expected, set(d) ^ expected


def test_t_d2_write_read_same_pipeline(monkeypatch):
    """T-D2：写入所得决定与读路径派生逐字段一致（写读同流水线）。"""
    rid = _mk()
    _confirm(rid)
    stored = store_module.get(rid)
    written = next(q.get("decision") for q in stored["verify"]["questions"]
                   if q.get("id") == "q1")
    from app.services.evidence import derive_claims_view
    view, _w, _cw = derive_claims_view(stored)
    read = next(q.get("decision") for q in view["verify"]["questions"]
                if q.get("id") == "q1")
    assert written["claim_id"] == read["claim_id"]
    assert written["claim_content_hash"] == read["claim_content_hash"]
    assert written["evidence_ids"] == read["evidence_ids"]
    # 写路径不落派生键（consistency 只在响应层组装）
    assert "consistency" not in written


def test_t_d3_decision_never_pollutes_claim_hash(monkeypatch):
    """T-D3（自查项）：写决定前后对象 claim_content_hash 不变（指纹免疫）。"""
    rid = _mk()
    from app.services.evidence import derive_claims_view
    before_view, _w, _cw = derive_claims_view(store_module.get(rid))
    before = next(q.get("claim_content_hash") for q in before_view["verify"]["questions"]
                  if q.get("id") == "q1")
    _confirm(rid)
    after_view, _w2, _cw2 = derive_claims_view(store_module.get(rid))
    after = next(q.get("claim_content_hash") for q in after_view["verify"]["questions"]
                 if q.get("id") == "q1")
    assert before == after, "决定键不得污染主张指纹"


def test_old_rows_response_byte_stable_except_new_keys(monkeypatch):
    """旧行 API 兼容：无 decision 键 → 响应除新增键外逐字节一致。"""
    rid = _mk()
    import json as _json
    before = client.get(f"/api/review/{rid}").json()
    client.get(f"/api/review/{rid}")
    after = client.get(f"/api/review/{rid}").json()
    # evidence_registry.rebuilt_at 每次 GET 即时重建，剔除后比对（先剔再序列化）
    for b in (before, after):
        if isinstance(b.get("evidence_registry"), dict):
            b["evidence_registry"].pop("rebuilt_at", None)
    assert _json.dumps(after, sort_keys=True, ensure_ascii=False) ==         _json.dumps(before, sort_keys=True, ensure_ascii=False)
