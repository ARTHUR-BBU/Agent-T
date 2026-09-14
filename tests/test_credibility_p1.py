"""可信度 P1 batch1：F01 覆盖截断 / F02 追问摘句核验 / F04 追问竞态 /
F05 docx 表序 / F08 押金 unless 客体绑定。
"""
from __future__ import annotations

import io
import json
from pathlib import Path
from unittest.mock import patch

import pytest

from app.services import llm_ask, model_review, quality as quality_service, scorecard
from app.services.checklist import run_checklist
from app.services.extract import extract_text
from app.services.llm_budget import ReviewBudget
from app.services.scorecard_prompts import (
    MAX_CONTRACT_CHARS,
    build_review_plan,
    coverage_from_plan,
)

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures"


# ---------- F01：覆盖不得因截断后块数齐全而报 limited=false ----------

def test_f01_over_24k_truncation_never_reports_unlimited():
    """~75k 字截到 ≤4×6k 后，即使 4 块全审阅成功，limited 仍必须为 true。"""
    text = ("甲" * 100 + "\n") * 750  # ~76k
    assert len(text) > 4 * MAX_CONTRACT_CHARS
    chunks, meta = build_review_plan(text, None, max_segments=4)
    assert len(chunks) == 4
    assert meta["fully_covered"] is False
    assert meta["truncation_reason"] == "max_segments_clip"
    assert meta["chars_covered"] < meta["original_chars"]
    assert meta["unread_ranges"], "截断中段必须记入 unread_ranges"
    cov = coverage_from_plan(meta, chunks_reviewed=4)
    assert cov["limited"] is True
    assert cov["truncation_reason"] == "max_segments_clip"
    assert cov["chunks_reviewed"] == cov["chunks_total"] == 4


def test_f01_oversized_clause_hard_split_tracks_coverage():
    """单条款超 MAX_CONTRACT_CHARS 硬切后，块数封顶仍须保留截断原因。"""
    text = "超" * 25000
    chunks, meta = build_review_plan(text, None, max_segments=4)
    assert len(chunks) <= 4
    cov = coverage_from_plan(meta, chunks_reviewed=len(chunks))
    if meta["truncation_reason"]:
        assert cov["limited"] is True
        assert cov["chars_covered"] < cov["original_chars"]
    else:
        # 硬切后恰好 ≤max_segments 且无 clip：才允许全覆盖
        assert meta["fully_covered"] is True


def test_f01_quality_map_partial_fail_limited(monkeypatch):
    monkeypatch.setenv("QUALITY_ENABLED", "true")
    para = "双方确认本条款为格式条款，双方均已知悉并同意其全部内容，任何修改须经书面协商一致后生效。\n"
    text = "长合同\n" + para * 140
    state = {"n": 0}

    def chat(system: str, user: str) -> str:
        state["n"] += 1
        if state["n"] == 1:
            raise RuntimeError("chunk 1 down")
        if "一致性轮" in user:
            return json.dumps({"observations": []}, ensure_ascii=False)
        return json.dumps(
            {
                "observations": [
                    {
                        "dimension": "impact",
                        "title": "格式条款提示",
                        "quote": "双方确认本条款为格式条款，双方均已知悉并同意其全部内容，任何修改须经书面协商一致后生效。",
                        "comment": "格式条款提示偏弱，建议补充说明义务。",
                    }
                ]
            },
            ensure_ascii=False,
        )

    out = quality_service.run_quality(text=text, items=[], chat_fn=chat)
    assert out["coverage"]["limited"] is True
    assert out["coverage"]["chunks_reviewed"] < out["coverage"]["chunks_total"]
    assert out["coverage"].get("truncation_reason")


def test_f01_quality_map_all_fail_keeps_coverage(monkeypatch):
    monkeypatch.setenv("QUALITY_ENABLED", "true")
    para = "双方确认本条款为格式条款，双方均已知悉并同意其全部内容，任何修改须经书面协商一致后生效。\n"
    text = "长合同\n" + para * 140

    def chat(system: str, user: str) -> str:
        raise RuntimeError("all down")

    out = quality_service.run_quality(text=text, items=[], chat_fn=chat)
    assert out["available"] is False
    assert out["coverage"] is not None
    assert out["coverage"]["limited"] is True
    assert out["coverage"]["chunks_reviewed"] == 0


def test_f01_budget_exhaust_marks_limited(monkeypatch):
    monkeypatch.setenv("QUALITY_ENABLED", "true")
    para = "双方确认本条款为格式条款，双方均已知悉并同意其全部内容，任何修改须经书面协商一致后生效。\n"
    text = "长合同\n" + para * 140

    def chat(system: str, user: str) -> str:
        return json.dumps(
            {
                "observations": [
                    {
                        "dimension": "impact",
                        "title": "t",
                        "quote": "双方确认本条款为格式条款，双方均已知悉并同意其全部内容，任何修改须经书面协商一致后生效。",
                        "comment": "提示偏弱，建议补充。",
                    }
                ]
            },
            ensure_ascii=False,
        )

    out = quality_service.run_quality(
        text=text, items=[], chat_fn=chat, budget=ReviewBudget(1)
    )
    assert out["coverage"]["limited"] is True
    assert out["coverage"]["truncation_reason"] == "budget_exhausted"
    assert out["coverage"]["chunks_reviewed"] < out["coverage"]["chunks_total"]


def test_f01_model_review_map_fail_fallback_coverage_limited():
    """map 全败回退头尾采样时，scorecard.available 可为 true，但 coverage.limited 必须 true。"""
    text = (FIXTURES / "procurement_long_backtoback.txt").read_text(encoding="utf-8")
    items = [
        {"id": "payment", "name": "付款条款", "segment": "C", "status": "通过", "note": "", "hits": [], "quote": ""},
        {"id": "breach", "name": "违约责任", "segment": "D", "status": "需关注", "note": "x", "hits": [], "quote": ""},
    ]
    reduce_payload = json.dumps(
        {
            "scorecard": {
                "summary": "规则结果汇总。",
                "segments": [
                    {"key": "C", "score": 8, "comment": "付款条款按规则结果展开", "gap_item_ids": []},
                    {"key": "D", "score": 10, "comment": "违约责任条款按规则结果展开", "gap_item_ids": []},
                ],
            },
            "candidates": [],
        },
        ensure_ascii=False,
    )
    calls = []

    def chat(system: str, user: str) -> str:
        calls.append((system, user))
        if "分段阅读" in system:
            raise RuntimeError("map 挂了")
        return reduce_payload

    result = model_review.run_model_review(
        text=text, items=items, policies=[], category="procurement", chat_fn=chat
    )
    sc = result["scorecard"]
    assert sc["available"] is True
    assert sc.get("coverage") is not None
    assert sc["coverage"]["limited"] is True
    assert sc["coverage"]["truncation_reason"] == "map_fail_fallback"


# ---------- F02：追问摘句核验 ----------

def _ask_with_payload(monkeypatch, payload, *, contract, clause_context="", item=None):
    for k in ("DEEPSEEK_API_KEY", "ZHIPU_API_KEY", "GLM_API_KEY", "XAI_API_KEY", "GROK_API_KEY"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("ZHIPU_API_KEY", "zk-test")
    item = item or {
        "id": "pay",
        "name": "价款与支付",
        "status": "需关注",
        "note": "预付",
        "quote": "预付全款",
    }
    with patch("app.services.llm_ask._chat_zhipu", return_value=json.dumps(payload, ensure_ascii=False)):
        return llm_ask.ask_about_item(
            question="风险大吗？",
            item=item,
            contract_text=contract,
            policies=[],
            clause_context=clause_context,
        )


def test_f02_fabricated_quote_not_verified(monkeypatch):
    payload = {
        "风险等级": "高",
        "这条在查啥": "价款",
        "原文在哪": "「甲方应支付天价违约金壹亿元」",
        "问题是啥": "伪造摘句",
        "建议怎么改": "改成分期",
        "改写稿": "甲方分期支付。",
        "还想问": "",
    }
    out = _ask_with_payload(monkeypatch, payload, contract="甲方预付全款后乙方发货。")
    assert out["ok"] is True
    assert out["quote_verified"] is False
    assert out["answer"]["原文在哪"] == "未定位到原文"
    assert not out["answer"].get("改写稿"), "未核验不得保留可复制改写稿"


def test_f02_empty_quote_honest(monkeypatch):
    payload = {
        "风险等级": "中",
        "这条在查啥": "价款",
        "原文在哪": "",
        "问题是啥": "未引用",
        "建议怎么改": "",
        "改写稿": "伪造改写",
        "还想问": "",
    }
    out = _ask_with_payload(monkeypatch, payload, contract="甲方预付全款后乙方发货。")
    assert out["quote_verified"] is False
    assert out["answer"]["原文在哪"] == "未定位到原文"
    assert not out["answer"].get("改写稿")


def test_f02_partial_fabricated_rejected(monkeypatch):
    """真句缀伪造尾巴 → 非连续原文，必须拒绝。"""
    payload = {
        "风险等级": "高",
        "这条在查啥": "价款",
        "原文在哪": "预付全款并放弃一切抗辩权利",
        "问题是啥": "拼接伪造",
        "建议怎么改": "改",
        "改写稿": "改写",
        "还想问": "",
    }
    out = _ask_with_payload(monkeypatch, payload, contract="甲方预付全款后乙方发货。")
    assert out["quote_verified"] is False


def test_f02_other_clause_quote_rejected_when_context_present(monkeypatch):
    contract = (
        "第一条 付款：甲方预付全款后乙方发货。\n"
        "第二条 保密：乙方不得泄露甲方商业秘密。"
    )
    clause_ctx = "第二条 保密：乙方不得泄露甲方商业秘密。"
    payload = {
        "风险等级": "高",
        "这条在查啥": "保密",
        "原文在哪": "甲方预付全款后乙方发货",
        "问题是啥": "拿付款句冒充保密依据",
        "建议怎么改": "改",
        "改写稿": "改写稿内容",
        "还想问": "",
    }
    out = _ask_with_payload(
        monkeypatch,
        payload,
        contract=contract,
        clause_context=clause_ctx,
        item={"id": "sec", "name": "保密", "status": "需关注", "note": "", "quote": "商业秘密"},
    )
    assert out["quote_verified"] is False
    assert out["answer"]["原文在哪"] == "未定位到原文"


def test_f02_real_quote_still_verified(monkeypatch):
    payload = {
        "风险等级": "高",
        "这条在查啥": "价款",
        "原文在哪": "「预付全款」",
        "问题是啥": "预付过高",
        "建议怎么改": "改为验收后付",
        "改写稿": "验收合格后支付全款。",
        "还想问": "",
    }
    out = _ask_with_payload(monkeypatch, payload, contract="甲方预付全款后乙方发货。")
    assert out["quote_verified"] is True
    assert "预付全款" in out["answer"]["原文在哪"]
    assert out["answer"]["改写稿"]


# ---------- F04：追问竞态守卫 ----------

def test_f04_ask_response_guard_logic_mirror():
    """镜像 app.js isAskResponseCurrent：rid/item/req 任一漂移即丢弃。"""

    def is_ask_response_current(state, snapshot):
        if not snapshot:
            return False
        if state["reviewId"] != snapshot["reviewId"]:
            return False
        if not state.get("askItem") or state["askItem"]["id"] != snapshot["itemId"]:
            return False
        if state["askReqSeq"] != snapshot["reqId"]:
            return False
        return True

    state = {"reviewId": "r1", "askItem": {"id": "pay"}, "askReqSeq": 3}
    snap = {"reviewId": "r1", "itemId": "pay", "reqId": 3}
    assert is_ask_response_current(state, snap) is True
    assert is_ask_response_current({**state, "reviewId": "r2"}, snap) is False
    assert is_ask_response_current({**state, "askItem": {"id": "other"}}, snap) is False
    assert is_ask_response_current({**state, "askReqSeq": 4}, snap) is False


def test_f04_app_js_contains_stale_ask_guard():
    src = (ROOT / "app" / "static" / "app.js").read_text(encoding="utf-8")
    assert "askReqSeq" in src
    assert "isAskResponseCurrent" in src
    assert "snapshot.reviewId" in src
    assert "snapshot.itemId" in src
    assert "snapshot.reqId" in src


# ---------- F05：docx 表序 ----------

def test_f05_docx_payment_table_stays_with_clause():
    raw = (FIXTURES / "docx_table_order.docx").read_bytes()
    text = extract_text("docx_table_order.docx", raw)
    assert "首付款" in text and "第三条" in text
    assert text.index("第二条") < text.index("首付款") < text.index("第三条"), (
        "付款表必须夹在第二条与第三条之间，不得被甩到文末"
    )


def test_f05_old_paragraphs_then_tables_would_misorder():
    """对照：若先段落再表格，首付款会落到第三条之后——本测试钉住回归方向。"""
    from docx import Document

    raw = (FIXTURES / "docx_table_order.docx").read_bytes()
    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                t = cell.text.strip()
                if t:
                    parts.append(t)
    legacy = "\n".join(parts)
    # 旧序：第三条在首付款前
    assert legacy.index("第三条") < legacy.index("首付款")
    # 新序已在上一测断言正确


# ---------- F08：押金 unless 客体绑定 ----------

def test_f08_equipment_protection_does_not_exempt_deposit():
    text = "第一条 乙方违反保密义务时，押金不予退还。第二条 甲方不得没收乙方自有设备。"
    by_id = {i["id"]: i for i in run_checklist(text, "lease")["items"]}
    assert by_id["deposit"]["status"] == "需关注", (
        f"设备保护不得洗白押金没收，实际 {by_id['deposit']['status']}"
    )


def test_f08_true_deposit_protection_still_passes():
    text = (
        "出租方（甲方）：某某置业有限公司。承租方（乙方）：某某科技有限公司。\n"
        "押金2万元。租赁期满，甲方不得拒绝返还保证金，不得没收押金。\n"
        "争议向法院起诉。本合同适用中华人民共和国法律。双方签字并加盖公章。"
    )
    by_id = {i["id"]: i for i in run_checklist(text, "lease")["items"]}
    assert by_id["deposit"]["status"] == "通过"
