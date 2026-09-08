"""外部审计整改回归（2026-09-08，ChatGPT 体检发现项）。

钉死：解析 fail-closed、上传硬限制、品类 422、后台任务化、
SQLite 存储与过期、评分头尾采样、追问错误不泄露、措辞降调。
"""
from __future__ import annotations

import io
import os
import tempfile
import time
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.extract import ExtractionError, extract_text
from app.services.scorecard import MAX_CONTRACT_CHARS, _clip_for_scoring, tier_of
from app.services.store import ReviewStore

ROOT = Path(__file__).resolve().parents[1]
client = TestClient(app)


# ---------- 1. 解析 fail-closed（P0：假成功产出全错报告） ----------

def test_zip_binary_disguised_as_docx_rejected():
    """损坏 docx（实为 ZIP 字节流）不得解码成乱码继续审查——必须报错。"""
    garbage = b"PK\x03\x04" + bytes(range(256)) * 40
    with pytest.raises(ExtractionError):
        extract_text("broken.docx", garbage)


def test_binary_disguised_as_txt_rejected():
    with pytest.raises(ExtractionError):
        extract_text("fake.txt", bytes(range(256)) * 64)


def test_old_doc_ole_format_clear_error():
    """旧版 .doc（OLE2 容器）无 docling 时明确指路，不吐乱码。"""
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512
    with pytest.raises(ExtractionError, match="另存为"):
        extract_text("legacy.doc", ole)


def test_docx_table_only_content_extracted():
    """正文全在表格里的 docx 必须能提取（审计复现的漏检场景）。"""
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "甲方：某某置业有限公司，法定代表人：张三。"
    table.cell(0, 1).text = "乙方：某某科技有限公司，法定代表人：李四。"
    table.cell(1, 0).text = "租赁期限自2026年10月1日起至2027年9月30日止。"
    table.cell(1, 1).text = "月租金1万元，押金2万元，租期未满退租押金不予退还。"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("table.docx", buf.getvalue())
    assert "某某置业" in text and "押金不予退还" in text


def test_scanned_pdf_no_text_layer_rejected():
    """无文字层的 PDF（pypdf 提取为空）必须报错，不得当空合同审查。"""
    with pytest.raises(ExtractionError):
        extract_text("scanned.pdf", b"%PDF-1.4\n" + b"\x00" * 256)


def test_normal_text_still_extracts():
    text = extract_text("normal.txt", "甲乙双方约定租金每月一万元。".encode("utf-8"))
    assert "租金" in text


# ---------- 2. 上传硬限制 ----------

def test_upload_oversize_rejected_413():
    big = b"\xff" * (10 * 1024 * 1024 + 1)
    r = client.post(
        "/api/upload",
        files={"file": ("big.txt", big, "text/plain")},
        data={"category": "procurement"},
    )
    assert r.status_code == 413


def test_upload_bad_extension_rejected_400():
    r = client.post(
        "/api/upload",
        files={"file": ("evil.exe", b"MZ...", "application/octet-stream")},
        data={"category": "procurement"},
    )
    assert r.status_code == 400


def test_upload_unknown_category_422_not_silent_fallback():
    """未知品类必须 422，此前会静默回退采购清单产出错误审查。"""
    r = client.post(
        "/api/upload",
        files={"file": ("c.txt", "甲方乙方".encode(), "text/plain")},
        data={"category": "employment"},
    )
    assert r.status_code == 422
    assert "employment" in r.json()["detail"]


# ---------- 3. 上传后台任务化 ----------

def test_upload_returns_review_id_immediately():
    """上传必须秒回 review_id（不再同步等待模型），状态为 processing 或已完成。"""
    import time as _t

    path = ROOT / "fixtures" / "lease_sample.txt"
    t0 = _t.time()
    r = client.post(
        "/api/upload",
        files={"file": ("l.txt", path.read_bytes(), "text/plain")},
        data={"category": "lease"},
    )
    elapsed = _t.time() - t0
    assert r.status_code == 200
    rid = r.json()["review_id"]
    assert rid
    assert elapsed < 10, f"上传不应同步等模型，实际 {elapsed:.1f}s"
    # 等后台任务收尾，避免后台线程写库与测试进程退出竞态
    deadline = time.time() + 30
    while time.time() < deadline:
        body = client.get(f"/api/review/{rid}").json()
        if body["status"] in ("done", "error"):
            break
        time.sleep(0.05)
    assert body["status"] == "done"


# ---------- 4. SQLite 存储与过期 ----------

def test_store_roundtrip_and_update(tmp_path):
    st = ReviewStore(db_path=str(tmp_path / "t.db"), ttl_seconds=3600)
    rid = st.create(filename="a.txt", status="processing", items=[])
    assert st.get(rid)["filename"] == "a.txt"
    st.update(rid, status="done", items=[{"id": "x"}])
    row = st.get(rid)
    assert row["status"] == "done" and row["items"] == [{"id": "x"}]
    assert row["filename"] == "a.txt", "update 只合并传入字段"


def test_store_persists_across_instances(tmp_path):
    db = str(tmp_path / "p.db")
    st1 = ReviewStore(db_path=db, ttl_seconds=3600)
    rid = st1.create(filename="a.txt", status="done")
    st2 = ReviewStore(db_path=db, ttl_seconds=3600)
    assert st2.get(rid)["status"] == "done", "重启（新实例）后数据仍在"


def test_store_ttl_expires(tmp_path):
    st = ReviewStore(db_path=str(tmp_path / "ttl.db"), ttl_seconds=0.05)
    rid = st.create(filename="a.txt", status="done")
    assert st.get(rid) is not None
    time.sleep(0.1)
    st.create(filename="b.txt", status="done")  # 触发过期清理
    assert st.get(rid) is None, "过期记录应被清除"


# ---------- 5. 评分头尾采样 ----------

def test_clip_keeps_head_and_tail():
    """签署/落款在尾部：截断必须保留尾部信号（审计指出只取头的注释失实）。"""
    text = "开头主体信息。".ljust(10) + "甲" * (MAX_CONTRACT_CHARS + 2000) + "尾部签字盖章要件。"
    clipped = _clip_for_scoring(text)
    assert len(clipped) <= MAX_CONTRACT_CHARS + 20
    assert clipped.startswith("开头主体信息")
    assert clipped.endswith("尾部签字盖章要件。"), "尾部（签署区）必须保留"
    assert "中段截断" in clipped


def test_clip_short_text_untouched():
    assert _clip_for_scoring("短合同") == "短合同"


# ---------- 6. 追问错误不泄露内部细节 ----------

def test_ask_error_message_fixed_no_exc_leak(monkeypatch):
    from app.services import llm_ask

    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    monkeypatch.setattr(
        llm_ask, "_chat_deepseek",
        lambda *a, **kw: (_ for _ in ()).throw(RuntimeError("http://internal-endpoint secret")),
    )
    result = llm_ask.ask_about_item(
        question="风险大吗？",
        item={"id": "x", "name": "押金", "status": "需关注", "note": "", "quote": "押金不予退还"},
        contract_text="押金不予退还。",
        policies=[],
    )
    assert result["ok"] is False
    assert "internal-endpoint" not in (result.get("error") or ""), "异常详情不得泄给用户"
    assert "稍后重试" in result["error"]


# ---------- 7. 措辞降调 ----------

def test_top_tier_label_not_overclaiming():
    """90+ 档不得用「基本没毛病」这类接近背书的措辞（外部审计）。"""
    assert tier_of(95)["label"] == "未见实质风险"
    assert tier_of(95)["label"] != "基本没毛病"


def test_report_all_pass_wording_softened():
    from app.services.report import build_report_docx

    row = {
        "id": "x", "filename": "a.txt", "category_label": "租赁合同",
        "created_at": "2026-09-08 10:00", "status": "done",
        "items": [{"id": "s", "name": "主体", "status": "通过", "note": "", "quote": "",
                   "hits": [], "category_na": False, "tag_source": "rule",
                   "needs_confirm": False}],
        "scorecard": {}, "blind_candidates": [], "blind_skipped_messages": [],
        "blind_enabled": False, "policies": [], "error": None, "text": "",
    }
    from docx import Document as _Doc
    data = build_report_docx(row)
    doc = _Doc(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "全部适用项均通过" not in text
    assert "规则初筛未命中风险项" in text


# ---------- 8. 门禁 P2/P3 修复（2026-09-08 小智娘报告） ----------

def test_ask_question_too_long_rejected(monkeypatch):
    """追问无长度限制会放大模型费用（外部审计 P1）——超 500 字直接拒绝。"""
    from app.services import llm_ask

    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    result = llm_ask.ask_about_item(
        question="甲" * 501,
        item={"id": "x", "name": "押金", "status": "需关注", "note": "", "quote": "押金不予退还"},
        contract_text="押金不予退还。",
        policies=[],
    )
    assert result["ok"] is False
    assert "500" in result["error"]


def test_store_marks_stale_processing_on_restart(tmp_path):
    """服务重启后遗留 processing 行必须被标记失败，不能让用户挂到 TTL（P3-2）."""
    db = str(tmp_path / "stale.db")
    st1 = ReviewStore(db_path=db, ttl_seconds=3600)
    rid = st1.create(filename="a.txt", status="processing", items=[])
    st2 = ReviewStore(db_path=db, ttl_seconds=3600)  # 模拟重启
    row = st2.get(rid)
    assert row["status"] == "error"
    assert "重新上传" in (row.get("error") or "")


def test_ask_prompt_uses_head_tail_clip(monkeypatch):
    """追问链路与评分卡同向：合同文本走头尾采样（P3-7 一致性）."""
    from app.services import llm_ask

    captured = {}
    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    monkeypatch.setattr(
        llm_ask, "_chat_deepseek",
        lambda k, s, u: (captured.update(prompt=u) or "{}"),
    )
    long_contract = "开头主体。" + "甲" * (MAX_CONTRACT_CHARS + 1000) + "尾部签字盖章。"
    llm_ask.ask_about_item(
        question="风险大吗？",
        item={"id": "x", "name": "押金", "status": "需关注", "note": "", "quote": "押金不予退还"},
        contract_text=long_contract,
        policies=[],
    )
    prompt = captured["prompt"]
    assert "开头主体。" in prompt and "尾部签字盖章。" in prompt


def test_model_review_error_reason_no_exc_leak(monkeypatch):
    """评分卡 LLM 异常时 reason 只给固定码（肉饼门禁 P2-2）——exc 含供应商
    URL/响应体，会经 /api/review 返回并印进客户 docx 报告。"""
    from app.services import model_review

    monkeypatch.setenv("DEEPSEEK_API_KEY", "sk-test")
    monkeypatch.setattr(
        model_review,
        "_call_llm",
        lambda *a, **kw: (_ for _ in ()).throw(
            RuntimeError("https://api.deepseek.com secret-response-body")
        ),
    )
    items = [{"id": "x", "name": "押金", "status": "需关注", "note": "", "quote": "押金不予退还"}]
    result = model_review.run_model_review(
        text="押金不予退还。", items=items, policies=[], category="lease"
    )
    sc = result["scorecard"]
    assert sc["available"] is False
    assert sc["reason"] == "llm_error", f"reason 不得携带异常详情：{sc['reason']}"
    assert "deepseek.com" not in (sc["reason"] or "")
