# -*- coding: utf-8 -*-
"""A2 表格抽取验尸与补缺 · 抽取层正式回归断言（A1 验尸结论：抽取层无罪，但断言为零）。

用户核定（2026-10-09）：.docx 抽取器已会按文档顺序读一级段落+一级表格，
但自动测试主要覆盖旧 .doc——本文件补齐 .docx 表格抽取的正式回归防线。
**不改 extract.py**：锁住现有行为，防后人回归。

覆盖：段落/表格顺序交错（锁 F05）、cell 全文进入、合并单元格不丢失、
空白格不产出噪音、嵌套表边界显式化（当前读不到=已知边界，非覆盖缺口）、
金标证物生产入口重放。
"""
from __future__ import annotations

import io

import pytest

from app.services.extract import extract_text


def _docx_bytes(build) -> bytes:
    """build(doc) 组装文档 → docx 字节（内存，无外部 fixture 依赖）。"""
    from docx import Document

    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- 1. 段落/表格顺序交错（锁 F05：表格行不得甩到文末） ----------

def test_paragraph_table_interleaved_order() -> None:
    def build(doc) -> None:
        doc.add_paragraph("第一条 甲方义务")
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "付款节点"
        t.cell(0, 1).text = "验收后 10 日"
        doc.add_paragraph("第二条 乙方义务")

    text = extract_text("t.docx", _docx_bytes(build))
    i_first = text.index("第一条")
    i_table = text.index("付款节点")
    i_second = text.index("第二条")
    assert i_first < i_table < i_second, (
        f"表格行被甩到段落后（F05 回退）: {text!r}")


# ---------- 2. 表格 cell 全文进入抽取文本 ----------

def test_table_cell_text_fully_extracted() -> None:
    """金标关键词只存在于表格 cell 时必须可抽取（A1 food 证物形态）。"""

    def build(doc) -> None:
        t = doc.add_table(rows=1, cols=3)
        t.cell(0, 0).text = "序号"
        t.cell(0, 1).text = "材料名称"
        t.cell(0, 2).text = "技术及质量要求"

    text = extract_text("t.docx", _docx_bytes(build))
    for kw in ("序号", "材料名称", "技术及质量要求"):
        assert kw in text, f"表头 {kw!r} 丢失"


# ---------- 3. 合并单元格：文本不丢失（重复可接受） ----------

def test_merged_cell_text_not_lost() -> None:
    def build(doc) -> None:
        t = doc.add_table(rows=1, cols=3)
        # 跨 3 列合并的标题格
        merged = t.cell(0, 0).merge(t.cell(0, 1)).merge(t.cell(0, 2))
        merged.text = "合同标的及验收标准一览"

    text = extract_text("t.docx", _docx_bytes(build))
    assert "合同标的及验收标准一览" in text, "合并单元格文本丢失"


# ---------- 4. 空白格：不产出噪音、不中断同行内容连接 ----------

def test_blank_cells_do_not_break_row_joins() -> None:
    def build(doc) -> None:
        t = doc.add_table(rows=1, cols=3)
        t.cell(0, 0).text = "保密要求"
        t.cell(0, 1).text = ""  # 空白格
        t.cell(0, 2).text = "双方不得向第三方披露"

    text = extract_text("t.docx", _docx_bytes(build))
    assert "保密要求" in text and "双方不得向第三方披露" in text
    # 同行内容仍在同一行（tab 连接未被空白格断开成噪音行）
    row_line = next(ln for ln in text.split("\n") if "保密要求" in ln)
    assert "双方不得向第三方披露" in row_line


# ---------- 5. 嵌套表：显式化当前边界（读不到=已知行为） ----------

def test_nested_table_currently_not_extracted() -> None:
    """嵌套表内容当前抽取不到（python-docx Table.rows 不下钻，A1 实锤）。

    显式钉死该边界：现有 17 份语料无嵌套表（A1 全量扫描 0 个），故为
    已知边界而非覆盖缺口；嵌套表支持是独立批次，做之前先改本断言。
    """
    def build(doc) -> None:
        doc.add_paragraph("正文段落锚点")
        outer = doc.add_table(rows=1, cols=1)
        inner = outer.cell(0, 0).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "嵌套表内的保密条款"

    text = extract_text("t.docx", _docx_bytes(build))
    assert "嵌套表内的保密条款" not in text, (
        "嵌套表已可抽取——请同步更新本边界断言与 A1 验尸报告")


# ---------- 6. 金标证物生产入口重放（A1 验尸证物固化） ----------

def test_food_golden_evidence_material_names_in_tables() -> None:
    """food 证物：表头「材料名称」（金标标的语义域词）只存在于表格 cell，
    抽取层必须读到（A1 验尸：抽取层无罪，词表缺口另批修）。"""
    raw = open("fixtures-real/food-procurement-xinjiang-2025.docx", "rb").read()
    text = extract_text("food-procurement-xinjiang-2025.docx", raw)
    assert "材料名称" in text
    assert "技术及质量要求" in text


def test_work_contract_golden_evidence_confidentiality_label() -> None:
    """work 证物：第四条「…及保密要求：」在正文段落且冒号后留白（A1 验尸
    改判依据）；抽取层读到标签本体——空栏判定由规则层负责（已测）。"""
    raw = open("fixtures-real/work-contract-gf-2000.docx", "rb").read()
    text = extract_text("work-contract-gf-2000.docx", raw)
    assert "保密要求" in text


def test_real_corpus_no_nested_tables() -> None:
    """全 17 份真实语料嵌套表=0（A1 全量扫描结论）；若未来语料含嵌套表，
    本断言失败提醒先解决嵌套表抽取边界。"""
    import glob

    from docx import Document
    from docx.oxml.ns import qn

    nested_files = []
    for f in glob.glob("fixtures-real/*.docx"):
        doc = Document(f)
        for tbl in doc.element.body.iter(qn("w:tbl")):
            if tbl.findall(".//" + qn("w:tbl")):
                nested_files.append(f)
                break
    assert not nested_files, f"语料出现嵌套表（超出已验证边界）: {nested_files}"


@pytest.mark.parametrize("fname", [
    "fixtures-real/food-procurement-xinjiang-2025.docx",
    "fixtures-real/work-contract-gf-2000.docx",
    "fixtures-real/cement-sale-gf-2008.docx",
])
def test_real_documents_extract_nonempty(fname: str) -> None:
    """A1 三份证物生产入口重放：非空、非乱码（fail-closed 底线）。"""
    raw = open(fname, "rb").read()
    text = extract_text(fname, raw)
    assert text and len(text) > 200
