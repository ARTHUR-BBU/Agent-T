"""M4 审查报告导出测试。

覆盖两层：
- 单元：build_report_docx 的结构红线（免责句必备、补盲候选标注、全文不落盘、NA 折叠）
- API：GET /api/review/{id}/report 的 200/404/409 与下载头

无网络、无 Key；docx 用 python-docx 读回断言。
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.report import DEFAULT_DISCLAIMER, build_report_docx
from app.services.store import store

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "fixtures" / "procurement_sample.txt"

client = TestClient(app)


def _doc_text(data: bytes) -> str:
    """读回 docx 的全部可见文本（段落 + 表格单元格）。"""
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _full_row() -> dict:
    return {
        "id": "a3f2b1c9d0e4",
        "filename": "采购合同v3.docx",
        "category": "procurement",
        "category_label": "采购合同",
        "status": "done",
        "created_at": "2026-09-06 10:30",
        "text": "全文哨兵句：这一段是合同全文，绝不能整段写进报告。",
        "items": [
            {"id": "payment", "name": "付款条件", "status": "需关注",
             "note": "签约即付全款，无质保金。", "quote": "合同签订后三日内支付全款。",
             "hits": [], "category_na": False, "tag_source": "rule", "needs_confirm": False},
            {"id": "jurisdiction", "name": "管辖与争议", "status": "未找到",
             "note": "全文无争议解决条款。", "quote": "",
             "hits": [], "category_na": False, "tag_source": "rule", "needs_confirm": False},
            {"id": "subject", "name": "主体信息", "status": "通过",
             "note": "双方主体信息齐备。", "quote": "甲方：某某有限公司。",
             "hits": [], "category_na": False, "tag_source": "rule", "needs_confirm": False},
            {"id": "ip", "name": "知识产权", "status": "本类不适用",
             "note": "采购类不适用。", "quote": "",
             "hits": [], "category_na": True, "tag_source": "rule", "needs_confirm": False},
        ],
        "scorecard": {},
        "blind_candidates": [
            {"id": "signature", "name": "签署与印章", "status": "需关注",
             "note": "仅有盖章栏无签字栏。", "quote": "本协议一式两份，双方各执一份。",
             "hits": [], "tag_source": "blind", "needs_confirm": True,
             "named_by_scorecard": True},
        ],
        "blind_skipped_messages": ["保密: 缺少原文依据，已跳过"],
        "blind_skipped_reason": None,
        "blind_enabled": True,
        "policies": ["政策一：先验收后付款，预留 10% 质保金。"],
        "error": None,
    }


# ---------- 单元层：报告结构红线 ----------

def test_report_cover_and_default_disclaimer():
    """封面四要素 + 免责句必备（scorecard 没带时代码补默认句）."""
    text = _doc_text(build_report_docx(_full_row()))
    assert "合同审查报告" in text
    assert "采购合同v3.docx" in text
    assert "采购合同" in text
    assert "2026-09-06 10:30" in text
    assert "a3f2b1c9d0e4" in text
    assert DEFAULT_DISCLAIMER in text, "免责句是必备项，缺失时必须兜底"


def test_report_scorecard_available_shows_total_and_caps():
    """评分卡可用时展示总分/档位/封顶说明，并用卡内免责句替代默认句."""
    row = _full_row()
    row["scorecard"] = {
        "available": True,
        "reason": None,
        "total": 72,
        "tier": {"label": "有实质风险", "hint": "先改完再谈签的事"},
        "summary": "规则结果汇总评分 72 分。",
        "segments": [],
        "caps_applied": ["因存在【付款条件】需关注项，总分已按上限 89 封顶（失分不能互相抵扣）"],
        "disclaimer": "以上都是机器给的参考意见，签之前建议找懂行的人再看一眼。",
        "advisory_only": True,
    }
    text = _doc_text(build_report_docx(row))
    assert "参考评分：72/100（有实质风险）" in text
    assert "89 封顶" in text
    assert "以逐条规则结果为准" in text, "评分必须标注从属地位"
    assert "机器给的参考意见" in text
    assert DEFAULT_DISCLAIMER not in text, "卡内自带免责句时不再重复默认句"


def test_report_scorecard_unavailable_notes_reason():
    """评分卡未生成：说明原因，不装作有分，规则结果照常输出."""
    row = _full_row()
    row["scorecard"] = {"available": False, "reason": "no_llm_key"}
    text = _doc_text(build_report_docx(row))
    assert "参考评分未生成（no_llm_key）" in text
    assert "参考评分：" not in text


def test_report_blind_candidates_marked_needs_confirm():
    """补盲候选单独成节 + 「需人工确认」标注；评分卡点名要带来源."""
    text = _doc_text(build_report_docx(_full_row()))
    assert "模型补盲候选" in text
    assert "需人工确认" in text
    assert "未经规则引擎确认" in text
    assert "来自评分卡点名" in text
    assert "仅有盖章栏无签字栏" in text


def test_report_free_text_fields_scrub_forbidden():
    """禁语纵深防御（小智娘 xz2a/2b 钉住的 P3 → 已修转正）：
    summary/条目 note/候选 note 里的禁语不得原样印进报告."""
    row = _full_row()
    row["scorecard"] = {
        "available": True, "reason": None, "total": 95,
        "tier": {"label": "基本可控", "hint": ""},
        "summary": "本合同没有问题，可以放心签署。",
        "segments": [], "caps_applied": [],
        "disclaimer": "", "advisory_only": True,
    }
    row["items"][0]["note"] = "虽有提示但可以忽略。"
    row["blind_candidates"][0]["note"] = "问题不大。"
    text = _doc_text(build_report_docx(row))
    for banned in ("本合同没有问题", "可以放心签署", "虽有提示但可以忽略", "问题不大"):
        assert banned not in text, f"禁语「{banned}」经报告回显"
    assert "【已过滤】" in text, "清洗必须留痕，不能静默删句"
    assert DEFAULT_DISCLAIMER in text, "卡内 disclaimer 被清空时落默认句"


def test_report_never_contains_full_text():
    """合同全文不写入报告——只允许原文摘句（泄露面控制）."""
    text = _doc_text(build_report_docx(_full_row()))
    assert "绝不能整段写进报告" not in text
    assert "合同签订后三日内支付全款" in text, "摘句应保留"


def test_report_na_items_folded_into_appendix():
    """本类不适用：不进逐条明细，折叠到附注."""
    text = _doc_text(build_report_docx(_full_row()))
    assert "【本类不适用】" not in text, "NA 项不得出现在逐条明细标题里"
    assert "本类不适用（未核查）：知识产权" in text


def test_report_attention_summary_table():
    """需关注/未找到进汇总表；全过时给出「均通过」."""
    text = _doc_text(build_report_docx(_full_row()))
    assert "付款条件" in text and "需关注" in text
    assert "管辖与争议" in text and "未找到" in text

    row = _full_row()
    row["items"] = [
        {"id": "s", "name": "主体信息", "status": "通过", "note": "", "quote": "",
         "hits": [], "category_na": False, "tag_source": "rule", "needs_confirm": False},
    ]
    assert "全部适用项均通过" in _doc_text(build_report_docx(row))


def test_report_policies_and_skip_messages():
    """政策摘句列出；补盲跳过原因进附注（不伪造候选行）."""
    text = _doc_text(build_report_docx(_full_row()))
    assert "先验收后付款" in text
    assert "保密: 缺少原文依据，已跳过" in text


def test_report_minimal_row_no_crash():
    """极端缺字段的行（老数据/异常态）不炸，兜底文案顶上."""
    text = _doc_text(build_report_docx({"id": "x1", "status": "done"}))
    assert "合同审查报告" in text
    assert "未命名" in text
    assert "未记录" in text
    assert DEFAULT_DISCLAIMER in text


def test_report_control_characters_stripped_not_fatal():
    """肉饼 P1：quote/note/文件名含控制字符（\x0b\x0c 等）不得让导出 500，
    也不得把控制字符写进 docx（lxml 会炸）——清洗后照常出报告."""
    row = _full_row()
    row["filename"] = "采购\x0c合同v3.docx"
    row["items"][0]["note"] = "签约即付全款\x0b，无质保金。"
    row["items"][0]["quote"] = "合同签订后\x00三日内支付全款。"
    row["blind_candidates"][0]["note"] = "仅有盖章栏\x1f无签字栏。"
    row["policies"][0] = "政策一：先验收后付款\x7f。"
    data = build_report_docx(row)  # 不抛即通过第一关
    text = _doc_text(data)
    assert "签约即付全款，无质保金" in text, "清洗只删控制字符，正文要保留"
    assert "合同签订后三日内支付全款" in text
    assert "仅有盖章栏无签字栏" in text
    for ch in ("\x00", "\x0b", "\x0c", "\x1f", "\x7f"):
        assert ch not in text


# ---------- API 层：下载端点 ----------

def test_api_report_download_after_upload():
    files = {"file": ("procurement_sample.txt", FIXTURE.read_bytes(), "text/plain")}
    rid = client.post("/api/upload", files=files, data={"category": "procurement"}).json()[
        "review_id"
    ]
    r = client.get(f"/api/review/{rid}/report")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in r.headers["content-disposition"]
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    assert r.content[:2] == b"PK", "docx 本质是 zip 包"

    # 报告与页面同源：文件名来自 store
    text = _doc_text(r.content)
    assert "合同审查报告" in text
    assert "procurement_sample.txt" in text
    assert rid in text


def test_api_report_unknown_id_404():
    assert client.get("/api/review/no-such-id/report").status_code == 404


def test_api_report_not_done_409():
    """审查未完成（processing/error）不能导出半成品报告."""
    rid = store.create(filename="x.txt", category="procurement", status="processing")
    r = client.get(f"/api/review/{rid}/report")
    assert r.status_code == 409

    rid2 = store.create(filename="x.txt", category="procurement", status="error", error="解析失败")
    assert client.get(f"/api/review/{rid2}/report").status_code == 409


def test_api_report_500_returns_fixed_message(monkeypatch):
    """肉饼 P2-1：生成器内部异常时 detail 必须是固定文案，异常串（可能含
    服务器路径/实现细节）只进日志不回客户端."""
    from app.api import routes as routes_module

    def boom(_row):
        raise OSError("/srv/secret/path 缺盘")

    monkeypatch.setattr(routes_module.report_service, "build_report_docx", boom)
    rid = store.create(filename="x.txt", category="procurement", status="done")
    r = client.get(f"/api/review/{rid}/report")
    assert r.status_code == 500
    assert r.json()["detail"] == "报告生成失败，请稍后重试"
    assert "/srv/secret/path" not in r.text
