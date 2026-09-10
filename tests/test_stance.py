"""立场输入测试（阶段 1.3，老钱矩阵全覆盖 + pc11 四层验收）。

老钱裁决书：立场=声明+防错配；立场错配永不构成拦截事由；不可审立场
不渲染（API 422 兜底）；报告声明必须含「核查口径」一致表述。
"""
from __future__ import annotations

import io
import json

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services import stance as stance_service
from app.services.checklist import run_checklist

client = TestClient(app)

ROOT = __import__("pathlib").Path(__file__).resolve().parents[1]
PC11 = ROOT / "fixtures" / "precheck" / "pc11_sale_seller_view.txt"
PROC = ROOT / "fixtures" / "procurement_sample.txt"

# 预审 stub：检出「销售合同（卖方视角）」——按新视角规则应可审（suggested=procurement）
_SELLER_VIEW = json.dumps(
    {
        "detected_type": "销售合同（卖方视角）",
        "is_supported": True,
        "suggested_category": "procurement",
        "confidence": "high",
        "summary": "条款密集偏向卖方保护",
    },
    ensure_ascii=False,
)


def _upload(category: str, stance: str | None = None, path=PROC, force: bool = False):
    data = {"category": category}
    if stance is not None:
        data["stance"] = stance
    if force:
        data["force"] = "1"
    return client.post("/api/upload", files={"file": (path.name, path.read_bytes(), "text/plain")}, data=data)


def _wait_done(rid: str, timeout: float = 10.0) -> dict:
    import time

    deadline = time.time() + timeout
    while time.time() < deadline:
        d = client.get(f"/api/review/{rid}").json()
        if d["status"] in ("done", "error"):
            return d
        time.sleep(0.1)
    raise AssertionError("审查超时未完成")


@pytest.fixture(autouse=True)
def _precheck_off(monkeypatch):
    # conftest 已关；此处显式再关一次防本文件内 fixture 顺序意外
    monkeypatch.setenv("PRECHECK_ENABLED", "false")


# ---------- stance 模块单测 ----------

def test_declaration_matrix_all_seven_combos():
    combos = {
        ("procurement", "buyer"): "核查口径",
        ("procurement", "neutral"): "默认按买方",
        ("lease", "lessee"): "核查口径",
        ("lease", "neutral"): "默认按承租方",
        ("nda", "disclosing"): "披露",
        ("nda", "receiving"): "接收",
        ("nda", "neutral"): "不预设立场",
    }
    for (cat, st), keyword in combos.items():
        text = stance_service.declaration(cat, st)
        assert keyword in text, f"{cat}/{st} 声明缺少关键语义项"
    # 措辞红线（老钱 Q4）：声明立场的版本必须含核查口径一致表述
    for cat, st in (("procurement", "buyer"), ("lease", "lessee"), ("nda", "disclosing"), ("nda", "receiving")):
        assert "核查口径" in stance_service.declaration(cat, st)


def test_declaration_unknown_combo_falls_back_neutral():
    assert stance_service.declaration("procurement", "ghost") == stance_service.declaration("procurement", "neutral")


def test_counterparty_marker_matrix():
    assert stance_service.has_counterparty_view_marker("procurement", "销售合同（卖方视角）")
    assert stance_service.has_counterparty_view_marker("lease", "租赁合同（出租方视角）")
    assert not stance_service.has_counterparty_view_marker("procurement", "采购合同")
    assert not stance_service.has_counterparty_view_marker("nda", "保密协议")
    # 分支 C 文案三品类齐备（NDA 无对方视角 → 空文案）
    assert "请勿直接采信" in stance_service.counterparty_view_notice("procurement")
    assert "请勿直接采信" in stance_service.counterparty_view_notice("lease")
    assert stance_service.counterparty_view_notice("nda") == ""


# ---------- /api/categories 元数据 ----------

def test_categories_expose_stances_matrix():
    cats = {c["id"]: c["stances"] for c in client.get("/api/categories").json()["categories"]}
    assert cats["procurement"]["view"] == "buyer"
    assert cats["procurement"]["allowed"] == ["neutral", "buyer"]
    assert cats["lease"]["allowed"] == ["neutral", "lessee"]
    # 老钱矩阵：NDA 双向可审
    assert cats["nda"]["allowed"] == ["neutral", "disclosing", "receiving"]
    assert cats["nda"]["labels"]["disclosing"] == "披露方"


# ---------- 422 兜底与放行矩阵 ----------

@pytest.mark.parametrize(
    "category, stance",
    [
        ("procurement", "supplier"),  # 卖方：采购不可审（UI 不渲染，API 兜底）
        ("lease", "lessor"),  # 出租方：租赁不可审
        ("procurement", "ghost"),  # 非法值
    ],
)
def test_disallowed_stance_rejected_422(category, stance):
    assert _upload(category, stance).status_code == 422


@pytest.mark.parametrize(
    "category, stance",
    [("nda", "disclosing"), ("nda", "receiving"), ("procurement", "buyer"), ("lease", "lessee")],
)
def test_allowed_stance_proceeds(category, stance):
    r = _upload(category, stance)
    assert r.status_code == 200, r.text
    assert r.json()["review_id"]


def test_default_stance_is_neutral():
    """不传 stance = neutral（所有现存调用零影响）。"""
    r = _upload("procurement")
    assert r.status_code == 200
    d = _wait_done(r.json()["review_id"])
    assert d["stance"] == "neutral"
    assert "默认按买方" in d["stance_declaration"]


# ---------- pc11 四层验收（老钱裁决第五节） ----------

def _seller_view_precheck(monkeypatch):
    monkeypatch.setenv("PRECHECK_ENABLED", "true")
    monkeypatch.setattr(
        "app.api.routes.precheck_service.run_precheck",
        lambda t, c, chat_fn=None, budget=None, stance="neutral": __import__(
            "app.services.precheck", fromlist=["PrecheckOutcome"]
        ).PrecheckOutcome(
            performed=True,
            result=__import__("app.services.precheck", fromlist=["PrecheckResult"]).PrecheckResult(
                **json.loads(_SELLER_VIEW)
            ),
        ),
    )


def test_pc11_layer1_buyer_stance_proceeds_without_dialog(monkeypatch):
    """层1 路由：pc11 + 采购 + 买方 → 不弹不支持/不弹品类确认，直接开审。"""
    _seller_view_precheck(monkeypatch)
    r = _upload("procurement", "buyer", path=PC11)
    assert r.status_code == 200
    body = r.json()
    assert body["status"] == "uploaded", "买方立场 + 卖方视角文件必须直接开审"
    assert body["review_id"]


def test_pc11_layer4_neutral_proceeds_with_notice(monkeypatch):
    """层4 对抗面：pc11 + 中性 → 照审 + 分支 C 知情提示（不阻断）。"""
    _seller_view_precheck(monkeypatch)
    r = _upload("procurement", "neutral", path=PC11)
    assert r.status_code == 200
    rid = r.json()["review_id"]
    d = _wait_done(rid)
    assert d["status"] == "done"
    assert d["precheck"]["stance_notice"] is True
    assert "请勿直接采信" in d["precheck"]["stance_notice_text"]
    assert "默认按买方" in d["stance_declaration"]


def test_pc11_layer2_rule_gold_hits_directional_clusters():
    """层2 规则金标：卖方保护条款在买方尺子下大面命中需关注。"""
    r = run_checklist(PC11.read_text(encoding="utf-8"), "procurement")
    by_id = {it["id"]: it["status"] for it in r["items"]}
    for item_id in ("subject", "payment", "unfair_terms", "governing_law", "signature"):
        assert by_id[item_id] == "需关注", f"pc11 金标：{item_id} 必须命中需关注"
    # hardline 双钉：逾期视为合格 / 定金没收 不得被 pass 词表洗成通过
    unfair = next(it for it in r["items"] if it["id"] == "unfair_terms")
    assert "逾期视为合格" in (unfair["quote"] or "") or any(
        "视为合格" in h for h in unfair["hits"]
    ), "「逾期视为合格」必须被 unfair_terms 直捕"


def test_pc11_buyer_stance_full_flow_declaration(monkeypatch):
    """层3 声明：完整审查后报告头出现买方声明（含核查口径一致表述）。"""
    _seller_view_precheck(monkeypatch)
    r = _upload("procurement", "buyer", path=PC11)
    d = _wait_done(r.json()["review_id"])
    assert "核查口径" in d["stance_declaration"]
    assert "买方" in d["stance_declaration"]

    report = client.get(f"/api/review/{d['id']}/report")
    assert report.status_code == 200
    doc = Document(io.BytesIO(report.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "审查立场：买方（采购方）" in text
    assert "核查口径" in text
    assert "品类知情提示" not in text, "买方立场下分支 C 不得出现（分支 B 静默）"


# ---------- docx 声明落库 ----------

def test_docx_neutral_declaration_and_notice_line(monkeypatch):
    _seller_view_precheck(monkeypatch)
    r = _upload("procurement", "neutral", path=PC11)
    d = _wait_done(r.json()["review_id"])
    report = client.get(f"/api/review/{d['id']}/report")
    doc = Document(io.BytesIO(report.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "审查立场：中性（未声明）" in text
    assert "品类知情提示" in text and "请勿直接采信" in text


def test_stance_notice_absent_without_precheck():
    """预审降级（无 Key/关闭）→ 无知情提示，但声明仍在（声明不依赖预审）。"""
    r = _upload("procurement", "neutral")
    d = _wait_done(r.json()["review_id"])
    assert d["precheck"] is None or d["precheck"].get("stance_notice") is False
    assert "默认按买方" in d["stance_declaration"]
